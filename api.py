import csv
from datetime import date
from io import BytesIO, StringIO
from pathlib import Path

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
from pydantic import BaseModel

from conexao import conectar
from ofx_service import ler_ofx, salvar_movimentacoes

app = FastAPI(title="Projeto OFX API")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class Movimentacao(BaseModel):
    data: date
    valor: float
    descricao: str
    categoria: str


class Importacao(BaseModel):
    movimentacoes: list[Movimentacao]
    arquivo: str | None = None


@app.get("/api/health")
def health():
    try:
        banco = conectar()
        banco.close()
        return {"conectado": True}
    except Exception as erro:
        raise HTTPException(status_code=503, detail=f"Falha ao conectar ao banco: {erro}")


@app.post("/api/ofx/preview")
async def preview_ofx(file: UploadFile = File(...)):
    if not file.filename or not file.filename.lower().endswith(".ofx"):
        raise HTTPException(status_code=400, detail="Selecione um arquivo .ofx.")
    try:
        movimentos = ler_ofx(await file.read())
        entradas = sum((m["valor"] for m in movimentos if m["valor"] > 0), 0)
        saidas = sum((m["valor"] for m in movimentos if m["valor"] < 0), 0)
        categorias_despesas = list(dict.fromkeys(
            [m["categoria"] for m in movimentos if m["valor"] < 0] + ["Outros"]
        ))
        return {
            "arquivo": file.filename,
            "quantidade": len(movimentos),
            "entradas": entradas,
            "saidas": saidas,
            "saldo": entradas + saidas,
            "categorias_despesas": categorias_despesas,
            "movimentacoes": [
                {**m, "data": m["data"].isoformat(), "valor": float(m["valor"])}
                for m in movimentos
            ],
        }
    except Exception as erro:
        raise HTTPException(status_code=400, detail=f"Não foi possível ler o OFX: {erro}")


@app.post("/api/ofx/save")
def save_ofx(importacao: Importacao):
    try:
        salvar_movimentacoes([m.model_dump() for m in importacao.movimentacoes])
        return {"salvo": True, "quantidade": len(importacao.movimentacoes)}
    except Exception as erro:
        raise HTTPException(status_code=500, detail=f"Falha ao salvar no banco: {erro}")


def _nome_arquivo(importacao, extensao):
    base = Path(importacao.arquivo or "movimentacoes").stem
    seguro = "".join(char for char in base if char.isascii() and (char.isalnum() or char in ("-", "_", " "))).strip()
    return f"{seguro or 'movimentacoes'}.{extensao}"


def _linhas_exportacao(movimentacoes):
    return [
        [m.data, m.descricao, m.categoria, "Entrada" if m.valor >= 0 else "Saída", m.valor]
        for m in movimentacoes
    ]


@app.post("/api/ofx/export/csv")
def export_csv(importacao: Importacao):
    conteudo = StringIO(newline="")
    escritor = csv.writer(conteudo, delimiter=";")
    escritor.writerow(["Data", "Descrição", "Categoria", "Tipo", "Valor (R$)"])
    for data, descricao, categoria, tipo, valor in _linhas_exportacao(importacao.movimentacoes):
        escritor.writerow([data.strftime("%d/%m/%Y"), descricao, categoria, tipo, f"{valor:.2f}".replace(".", ",")])
    return Response(
        content="\ufeff" + conteudo.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{_nome_arquivo(importacao, "csv")}"'},
    )


@app.post("/api/ofx/export/excel")
def export_excel(importacao: Importacao):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Movimentações"
        sheet.append(["Data", "Descrição", "Categoria", "Tipo", "Valor (R$)"])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="347B5D")
        for data, descricao, categoria, tipo, valor in _linhas_exportacao(importacao.movimentacoes):
            sheet.append([data, descricao, categoria, tipo, valor])
            sheet.cell(sheet.max_row, 1).number_format = "DD/MM/YYYY"
            sheet.cell(sheet.max_row, 5).number_format = '"R$" #,##0.00'
        sheet.column_dimensions["A"].width = 14
        sheet.column_dimensions["B"].width = 38
        sheet.column_dimensions["C"].width = 30
        sheet.column_dimensions["D"].width = 14
        sheet.column_dimensions["E"].width = 18
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        arquivo = BytesIO()
        workbook.save(arquivo)
        return Response(
            content=arquivo.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{_nome_arquivo(importacao, "xlsx")}"'},
        )
    except Exception as erro:
        raise HTTPException(status_code=500, detail=f"Não foi possível gerar o Excel: {erro}")

